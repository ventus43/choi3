"""/reports  (CHOIREPORT) — 주일예배 구역별 자동분석(7Ius67Cp) 결과 저장.

본프로젝트(choi3 백오피스) 로그인과 완전히 분리된 별도 인증 영역이다.
X-Office-Auth 토큰으로는 이 아래 라우트에 접근할 수 없고, 반대로 여기서 발급한
X-Report-Auth 토큰으로는 /outreach, /meetings 등 다른 라우트에 접근할 수 없다
(auth.require_auth 의 EXEMPT_PREFIXES 참고).

이 안에 두 종류의 토큰이 공존한다 — 둘 다 X-Report-Auth 헤더, 같은 서명키를 쓰지만 용도가
다르다:
  - 부서 토큰(g.dept_id) : /reports/auth/login (부서 비밀번호) 으로 발급. 그 부서의
    구역·인원·주간보고 데이터만 보고/저장할 수 있다.
  - 관리자 토큰(g.is_report_admin) : /reports/admin/login (CHOIDEPT_ADMIN 비밀번호) 으로
    발급. 부서(CHOIDEPT) 자체를 추가/수정/삭제할 수 있지만, 특정 부서 데이터는 못 본다.
"""
import json
from io import BytesIO
from zipfile import ZipFile, BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

import pymysql
from flask import Blueprint, g, jsonify, request, send_file

from auth import issue_report_admin_token, issue_report_token, require_report_auth
from config import REPORT_SESSION_TTL
from db import db_cursor
from sqlpatch import build_update

reports_bp = Blueprint('reports', __name__)
reports_bp.before_request(require_report_auth)


def _require_dept():
    if g.dept_id is None:
        return jsonify({'message': '부서 로그인이 필요합니다.'}), 403
    return None


def _require_admin():
    if not g.is_report_admin:
        return jsonify({'message': '관리자 권한이 필요합니다.'}), 403
    return None


@reports_bp.route('/reports/auth/login', methods=['POST'])
def reports_login():
    body = request.get_json(silent=True) or {}
    with db_cursor() as cur:
        cur.execute('SELECT ID FROM CHOIDEPT WHERE PASSWORD=%s', (body.get('password'),))
        row = cur.fetchone()
    if not row:
        return jsonify({'message': '비밀번호가 올바르지 않습니다.'}), 401
    return jsonify({'token': issue_report_token(row['ID']), 'ttl': REPORT_SESSION_TTL})


@reports_bp.route('/reports/profile', methods=['GET'])
def report_profile():
    if (err := _require_dept()):
        return err
    with db_cursor() as cur:
        cur.execute('SELECT NAME FROM CHOIDEPT WHERE ID=%s', (g.dept_id,))
        dept = cur.fetchone()
    if not dept:
        return jsonify({'message': '부서를 찾을 수 없습니다.'}), 404
    return jsonify({'name': dept['NAME']})


@reports_bp.route('/reports/roster', methods=['GET'])
def report_roster():
    """구역명단 대조(병합 시 이름 검증·대시보드)·"인원 관리" 탭 목록용 — CHOIMEMBER 전체
    필드가 아니라 NTT_ID(수정/삭제용 식별자)·이름·구역·사명여부만 내려준다(report 토큰은
    최소 권한이라 TA 등 다른 필드는 제외). 동명이인(같은 구역에 이름이 같은 인원)이 실제로
    있어, 구역+이름만으론 서로 다른 사람을 구분할 수 없다 — ISMISSION 까지 포함해 최대한
    구분한다. 로그인한 부서 소속 인원만 내려준다."""
    if (err := _require_dept()):
        return err
    with db_cursor() as cur:
        cur.execute('SELECT NTT_ID, NAME, GU, ISMISSION FROM CHOIMEMBER WHERE DEPT_ID=%s', (g.dept_id,))
        return jsonify(cur.fetchall())


@reports_bp.route('/reports/member', methods=['POST'])
def report_member_create():
    """부서 담당자가 "인원 관리" 메뉴에서 직접 인원을 등록 — choi3 메인 SPA(구역 관리)와
    동일한 CHOIMEMBER 테이블을 공유하며, 로그인한 부서 소속으로 추가된다."""
    if (err := _require_dept()):
        return err
    body = request.get_json(silent=True) or {}
    name = (body.get('NAME') or '').strip()
    if not name:
        return jsonify({'message': '이름은 필수 입력입니다.'}), 400
    with db_cursor(commit=True) as cur:
        # NTT_ID 는 클라이언트가 안 보냄 — DB 에서 MAX+1 로 채번 (members.py 의 member_create 와 동일 방식)
        cur.execute('SELECT COALESCE(MAX(NTT_ID), 0) + 1 AS n FROM CHOIMEMBER')
        new_id = cur.fetchone()['n']
        cur.execute(
            'INSERT INTO CHOIMEMBER (NTT_ID, DEPT_ID, GU, NAME, TA, ISMISSION) VALUES (%s,%s,%s,%s,%s,%s)',
            (new_id, g.dept_id, body.get('GU') or None, name, body.get('TA', 0), body.get('ISMISSION', 'N')),
        )
        cur.execute('SELECT NTT_ID, NAME, GU, ISMISSION FROM CHOIMEMBER WHERE NTT_ID=%s', (new_id,))
        return jsonify(cur.fetchone()), 201


MAX_IMPORT_ROWS = 1000


def validate_import_rows(rows):
    """Preview and commit use identical validation; never trust browser validation."""
    if not isinstance(rows, list) or not rows or len(rows) > MAX_IMPORT_ROWS:
        raise ValueError('명단은 1명 이상, 최대 1,000명까지 등록할 수 있습니다.')
    clean, errors = [], []
    for index, row in enumerate(rows, 2):
        if not isinstance(row, dict):
            errors.append(f'{index}행: 구역과 이름을 확인해 주세요.')
            continue
        name, gu = row.get('NAME'), row.get('GU')
        name = name.strip() if isinstance(name, str) else ''
        try:
            zone = float(gu)
            valid = not isinstance(gu, bool) and zone.is_integer() and 1 <= zone <= 2147483647
        except (TypeError, ValueError, OverflowError):
            valid = False
        if not name or len(name) > 20 or name.startswith('=') or not valid:
            errors.append(f'{row.get("row", index)}행: 구역은 양의 정수, 이름은 1~20자로 입력해 주세요. 수식은 사용할 수 없습니다.')
        else:
            clean.append({'GU': int(zone), 'NAME': name, 'row': row.get('row', index)})
    return clean, errors


def classify_import_rows(rows, existing):
    seen = {(str(p['GU']), p['NAME']) for p in existing}
    result = []
    for row in rows:
        key = (str(row['GU']), row['NAME'])
        duplicate = key in seen
        result.append({**row, 'duplicate': duplicate})
        seen.add(key)
    return result


@reports_bp.route('/reports/member/template', methods=['GET'])
def report_member_template():
    if (err := _require_dept()):
        return err
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '명단'
    sheet.append(['구역', '이름'])
    example = workbook.create_sheet('작성 예시')
    for row in [['구역', '이름'], [1, '홍길동'], [1, '김민수'], [2, '이서연']]:
        example.append(row)
    example['D1'] = '입력 안내'
    example['D2'] = '명단 시트 2행부터 실제 인원을 입력하세요. 이 예시 시트는 등록하지 않습니다.'
    example['D3'] = '구역은 숫자만, 이름은 최대 20자. 한 행에 한 명, 최대 1,000명.'
    example['D4'] = '같은 구역·이름은 중복으로 제외합니다. 동명이인은 개별 등록하세요.'
    for ws in workbook:
        ws.freeze_panes = 'A2'
        ws.column_dimensions['A'].width = 14
        ws.column_dimensions['B'].width = 24
        for cell in ws[1][:2]:
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='0B6B68')
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name='인원등록_양식.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@reports_bp.route('/reports/member/import/preview', methods=['POST'])
def report_member_import_preview():
    if (err := _require_dept()):
        return err
    if request.content_length and request.content_length > 5 * 1024 * 1024:
        return jsonify({'message': '5MB 이하의 엑셀 파일을 선택해 주세요.'}), 400
    upload = request.files.get('file')
    if not upload or not (upload.filename or '').lower().endswith('.xlsx'):
        return jsonify({'message': '.xlsx 형식의 양식을 선택해 주세요.'}), 400
    data = upload.read(5 * 1024 * 1024 + 1)
    if len(data) > 5 * 1024 * 1024:
        return jsonify({'message': '파일이 너무 큽니다. 5MB 이하로 나눠 주세요.'}), 400
    try:
        with ZipFile(BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 20 * 1024 * 1024:
                raise ValueError('압축 해제 크기가 너무 큽니다. 제공된 양식에 값만 복사해 주세요.')
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = workbook['명단'] if '명단' in workbook.sheetnames else workbook.worksheets[0]
            if sheet.max_row and sheet.max_row > MAX_IMPORT_ROWS + 1:
                raise ValueError('명단 시트는 헤더를 포함해 1,001행 이내로 저장해 주세요.')
            sheet.reset_dimensions()
            iterator = sheet.iter_rows(max_row=MAX_IMPORT_ROWS + 2, max_col=2, values_only=True)
            if tuple(str(v or '').strip() for v in next(iterator)) != ('구역', '이름'):
                raise ValueError('첫 행 A1은 구역, B1은 이름이어야 합니다.')
            rows = [{'GU': gu, 'NAME': name, 'row': i} for i, (gu, name) in enumerate(iterator, 2)
                    if gu is not None or name is not None]
            if any(row['row'] > MAX_IMPORT_ROWS + 1 for row in rows):
                raise ValueError('명단 시트 2~1,001행에 최대 1,000명을 입력해 주세요.')
            clean, errors = validate_import_rows(rows)
        finally:
            workbook.close()
    except (ValueError, BadZipFile, KeyError, StopIteration) as exc:
        return jsonify({'message': str(exc) or '엑셀 양식의 내용을 확인해 주세요.'}), 400
    except Exception:
        return jsonify({'message': '엑셀 파일을 읽지 못했습니다. 제공된 양식에 값만 복사해 다시 저장해 주세요.'}), 400
    with db_cursor() as cur:
        cur.execute('SELECT GU, NAME FROM CHOIMEMBER WHERE DEPT_ID=%s', (g.dept_id,))
        result = classify_import_rows(clean, cur.fetchall())
    return jsonify({'rows': result, 'errors': errors})


@reports_bp.route('/reports/member/import', methods=['POST'])
def report_member_import():
    if (err := _require_dept()):
        return err
    body = request.get_json(silent=True) or {}
    if not isinstance(body, dict):
        return jsonify({'message': '명단 데이터를 확인해 주세요.'}), 400
    try:
        rows, errors = validate_import_rows(body.get('rows'))
    except ValueError as exc:
        return jsonify({'message': str(exc)}), 400
    if errors:
        return jsonify({'message': errors[0]}), 400
    try:
        with db_cursor(commit=True) as cur:
            # Serialize ID allocation/imports; lock existing rows and the upper index gap.
            cur.execute('SELECT NTT_ID, DEPT_ID, GU, NAME FROM CHOIMEMBER ORDER BY NTT_ID FOR UPDATE')
            all_members = cur.fetchall()
            classified = classify_import_rows(rows, [p for p in all_members if p['DEPT_ID'] == g.dept_id])
            new_rows = [p for p in classified if not p['duplicate']]
            next_id = max((p['NTT_ID'] for p in all_members), default=0) + 1
            if new_rows:
                cur.executemany('INSERT INTO CHOIMEMBER (NTT_ID, DEPT_ID, GU, NAME, TA, ISMISSION) VALUES (%s,%s,%s,%s,%s,%s)',
                                [(next_id+i, g.dept_id, p['GU'], p['NAME'], '0', 'N') for i,p in enumerate(new_rows)])
    except (pymysql.IntegrityError, pymysql.OperationalError):
        return jsonify({'message': '동시에 명단이 변경되어 등록하지 못했습니다. 명단을 다시 확인하고 재시도해 주세요.'}), 409
    return jsonify({'created': len(new_rows), 'skipped': len(rows)-len(new_rows)})


@reports_bp.route('/reports/member/<int:ntt_id>', methods=['PUT'])
def report_member_update(ntt_id):
    """"인원 관리" 탭의 수정 — 로그인한 부서 소속 인원만 고칠 수 있다(다른 부서 소유
    NTT_ID 를 추측해 넣어도 DEPT_ID 불일치로 막힘)."""
    if (err := _require_dept()):
        return err
    body = request.get_json(silent=True) or {}
    col_map = {'NAME': 'NAME', 'GU': 'GU', 'ISMISSION': 'ISMISSION'}
    sets, vals = build_update(col_map, body, blank_to_null=False)
    if not sets:
        return jsonify({'message': '변경 항목 없음'}), 400
    vals += [ntt_id, g.dept_id]
    with db_cursor(commit=True) as cur:
        cur.execute(
            f'UPDATE CHOIMEMBER SET {",".join(sets)} WHERE NTT_ID=%s AND DEPT_ID=%s',
            vals,
        )
        if cur.rowcount == 0:
            return jsonify({'message': '해당 인원을 찾을 수 없습니다.'}), 404
        return jsonify({'ok': True})


@reports_bp.route('/reports/member/<int:ntt_id>', methods=['DELETE'])
def report_member_delete(ntt_id):
    """"인원 관리" 탭의 삭제 — 로그인한 부서 소속 인원만 지울 수 있다."""
    if (err := _require_dept()):
        return err
    with db_cursor(commit=True) as cur:
        cur.execute('DELETE FROM CHOIMEMBER WHERE NTT_ID=%s AND DEPT_ID=%s', (ntt_id, g.dept_id))
        if cur.rowcount == 0:
            return jsonify({'message': '해당 인원을 찾을 수 없습니다.'}), 404
        return ('', 204)


@reports_bp.route('/reports/member/zone/<gu>', methods=['DELETE'])
def report_member_delete_zone(gu):
    """"인원 관리" 탭의 구역별 전체삭제 — 로그인한 부서 소속, 해당 구역 인원을 한 번에 지운다
    (개별 삭제와 동일하게 DEPT_ID로 소유권 확인, 다른 부서 구역은 건드릴 수 없음)."""
    if (err := _require_dept()):
        return err
    with db_cursor(commit=True) as cur:
        cur.execute('DELETE FROM CHOIMEMBER WHERE GU=%s AND DEPT_ID=%s', (gu, g.dept_id))
        return jsonify({'deleted': cur.rowcount})


@reports_bp.route('/reports', methods=['GET'])
def report_list():
    if (err := _require_dept()):
        return err
    with db_cursor() as cur:
        cur.execute('SELECT REPORT_DT FROM CHOIREPORT WHERE DEPT_ID=%s ORDER BY REPORT_DT DESC', (g.dept_id,))
        return jsonify([r['REPORT_DT'].strftime('%Y-%m-%d') for r in cur.fetchall()])


@reports_bp.route('/reports/<date_str>', methods=['GET'])
def report_get(date_str):
    if (err := _require_dept()):
        return err
    with db_cursor() as cur:
        cur.execute('SELECT DATA FROM CHOIREPORT WHERE DEPT_ID=%s AND REPORT_DT=%s', (g.dept_id, date_str))
        row = cur.fetchone()
        if not row:
            return jsonify({'message': '기록이 없습니다.'}), 404
        return jsonify(json.loads(row['DATA']))


@reports_bp.route('/reports/<date_str>', methods=['PUT'])
def report_put(date_str):
    if (err := _require_dept()):
        return err
    body = request.get_json(silent=True) or {}
    data = json.dumps(body, ensure_ascii=False)
    with db_cursor(commit=True) as cur:
        cur.execute(
            'INSERT INTO CHOIREPORT (DEPT_ID, REPORT_DT, DATA) VALUES (%s,%s,%s)'
            ' ON DUPLICATE KEY UPDATE DATA=%s',
            (g.dept_id, date_str, data, data),
        )
        return jsonify({'ok': True})


# ── 부서 관리자 화면(7Ius67Cp.html 안의 별도 "관리자" 탭) ──────────────────────
# CHOIDEPT_ADMIN 비밀번호로 로그인 → 부서(CHOIDEPT) 추가/수정/삭제. 특정 부서 데이터
# (/reports/roster, /reports/<date> 등)는 이 토큰으로 접근 불가 — _require_dept 가 막는다.

@reports_bp.route('/reports/admin/login', methods=['POST'])
def reports_admin_login():
    body = request.get_json(silent=True) or {}
    with db_cursor() as cur:
        cur.execute('SELECT PASSWORD FROM CHOIDEPT_ADMIN WHERE ID=1')
        row = cur.fetchone()
    if not row or body.get('password') != row['PASSWORD']:
        return jsonify({'message': '비밀번호가 올바르지 않습니다.'}), 401
    return jsonify({'token': issue_report_admin_token(), 'ttl': REPORT_SESSION_TTL})


@reports_bp.route('/reports/admin/depts', methods=['GET'])
def reports_admin_dept_list():
    if (err := _require_admin()):
        return err
    with db_cursor() as cur:
        cur.execute('SELECT ID, NAME, PASSWORD FROM CHOIDEPT ORDER BY ID')
        return jsonify(cur.fetchall())


@reports_bp.route('/reports/admin/depts', methods=['POST'])
def reports_admin_dept_create():
    if (err := _require_admin()):
        return err
    body = request.get_json(silent=True) or {}
    name = (body.get('NAME') or '').strip()
    password = (body.get('PASSWORD') or '').strip()
    if not name or not password:
        return jsonify({'message': '부서명과 비밀번호를 모두 입력해 주세요.'}), 400
    try:
        with db_cursor(commit=True) as cur:
            cur.execute('INSERT INTO CHOIDEPT (NAME, PASSWORD) VALUES (%s,%s)', (name, password))
            new_id = cur.lastrowid
            cur.execute('SELECT ID, NAME, PASSWORD FROM CHOIDEPT WHERE ID=%s', (new_id,))
            return jsonify(cur.fetchone()), 201
    except pymysql.err.IntegrityError:
        return jsonify({'message': '이미 사용 중인 비밀번호입니다.'}), 400


@reports_bp.route('/reports/admin/depts/<int:dept_id>', methods=['PUT'])
def reports_admin_dept_update(dept_id):
    if (err := _require_admin()):
        return err
    body = request.get_json(silent=True) or {}
    col_map = {'NAME': 'NAME', 'PASSWORD': 'PASSWORD'}
    sets, vals = build_update(col_map, body, blank_to_null=False)
    if not sets:
        return jsonify({'message': '변경 항목 없음'}), 400
    vals.append(dept_id)
    try:
        with db_cursor(commit=True) as cur:
            cur.execute(f'UPDATE CHOIDEPT SET {",".join(sets)} WHERE ID=%s', vals)
            return jsonify({'ok': True})
    except pymysql.err.IntegrityError:
        return jsonify({'message': '이미 사용 중인 비밀번호입니다.'}), 400


@reports_bp.route('/reports/admin/depts/<int:dept_id>', methods=['DELETE'])
def reports_admin_dept_delete(dept_id):
    if (err := _require_admin()):
        return err
    with db_cursor(commit=True) as cur:
        cur.execute('SELECT COUNT(*) AS n FROM CHOIMEMBER WHERE DEPT_ID=%s', (dept_id,))
        if cur.fetchone()['n'] > 0:
            return jsonify({'message': '소속 인원이 있어 삭제할 수 없습니다.'}), 400
        cur.execute('DELETE FROM CHOIDEPT WHERE ID=%s', (dept_id,))
        return ('', 204)
