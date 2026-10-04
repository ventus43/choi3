import io
import sys
import unittest
from pathlib import Path
from contextlib import contextmanager
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from flask import Flask, g
from openpyxl import Workbook, load_workbook
from routes import reports

class ImportTests(unittest.TestCase):
    def setUp(self):
        self.app = Flask(__name__)
        self.app.register_blueprint(reports.reports_bp)
        self.app.before_request_funcs['reports'] = [lambda: setattr(g, 'dept_id', 1)]
        self.client = self.app.test_client()
        self.inserted = []
        self.existing = [{'NTT_ID': 7, 'DEPT_ID': 1, 'GU': '1', 'NAME': '기존'}]
        owner = self
        class Cursor:
            def execute(self, *args): pass
            def fetchall(self): return owner.existing
            def executemany(self, sql, values): owner.inserted.extend(values)
        @contextmanager
        def cursor(**kwargs): yield Cursor()
        self.mock = patch.object(reports, 'db_cursor', cursor)
        self.mock.start()
        self.addCleanup(self.mock.stop)

    def upload(self, rows):
        wb = Workbook()
        wb.active.title = '명단'
        for row in rows: wb.active.append(row)
        file = io.BytesIO(); wb.save(file); file.seek(0)
        return self.client.post('/reports/member/import/preview', data={'file': (file, '명단.xlsx')})

    def test_template(self):
        response = self.client.get('/reports/member/template')
        self.assertEqual(response.status_code, 200)
        wb = load_workbook(io.BytesIO(response.data))
        self.assertEqual(list(wb['명단'].values), [('구역', '이름')])
        self.assertIn('작성 예시', wb.sheetnames)

    def test_preview_duplicates_and_errors(self):
        result = self.upload([['구역','이름'],[1,'기존'],[2,'새인원'],[2,'새인원'],[-1,'오류'],[3,'=A1']])
        self.assertEqual(result.status_code, 200)
        self.assertEqual([p['duplicate'] for p in result.json['rows']], [True,False,True])
        self.assertEqual(len(result.json['errors']), 2)
        self.assertEqual(self.inserted, [])

    def test_invalid_headers_and_empty(self):
        for rows in [[['이름','구역']], [['구역','이름']]]:
            self.assertEqual(self.upload(rows).status_code, 400)

    def test_commit_rechecks_duplicates_and_scope(self):
        response = self.client.post('/reports/member/import', json={'rows':[{'GU':1,'NAME':'기존'},{'GU':2,'NAME':'신규'},{'GU':2,'NAME':'신규'}]})
        self.assertEqual(response.json, {'created':1,'skipped':2})
        self.assertEqual(self.inserted, [(8,1,2,'신규','0','N')])

    def test_invalid_batch_writes_nothing(self):
        result=self.client.post('/reports/member/import',json={'rows':[{'GU':1,'NAME':'정상'},{'GU':0,'NAME':'오류'}]})
        self.assertEqual(result.status_code,400)
        self.assertEqual(self.inserted,[])

    def test_admin_denied(self):
        self.app.before_request_funcs['reports']=[lambda:setattr(g,'dept_id',None)]
        self.assertEqual(self.client.get('/reports/member/template').status_code,403)
        self.assertEqual(self.client.post('/reports/member/import',json={'rows':[]}).status_code,403)

if __name__ == '__main__': unittest.main()
